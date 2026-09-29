"""Smart Attendance System: face recognition (InsightFace) + vector search (FAISS).

Enrol students from a folder of photos, sample the classroom camera(s) at random
moments during a lecture, and mark a student Present when they are seen in enough
samples. Writes an Excel report with photos.

Usage:
    python attendance_system.py                                  # webcam 0, 1-hour lecture, 10 samples
    python attendance_system.py --source 0 --source 1            # two cameras (multi-camera fusion)
    python attendance_system.py --source lecture.mp4             # recorded lecture video
    python attendance_system.py --source demo/classroom.jpg --dataset demo/dataset --snapshots demo/output
"""
import argparse
import os
import random
import time
from datetime import datetime

import cv2
import faiss
import numpy as np
from insightface.app import FaceAnalysis
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill

EMBEDDING_DIM = 512
FRAME_WIDTH = 640
IMAGE_EXTS = (".jpg", ".jpeg", ".png", ".bmp", ".webp")


# =========================
# INITIALIZE MODEL
# =========================

def load_model(det_thresh=0.35):
    """Load InsightFace buffalo_l (detection + ArcFace embeddings); uses CUDA when available."""
    import onnxruntime

    print("🔄 Loading InsightFace model...")
    available = onnxruntime.get_available_providers()
    providers = [p for p in ("CUDAExecutionProvider", "CPUExecutionProvider") if p in available]
    app = FaceAnalysis(name="buffalo_l", providers=providers)
    app.prepare(ctx_id=0 if "CUDAExecutionProvider" in providers else -1, det_thresh=det_thresh)
    return app


# =========================
# BUILD FAISS DATABASE
# =========================

def build_database(face_app, dataset_path):
    """Embed every photo in dataset/<student>/ and add it to a cosine-similarity FAISS index."""
    print("📚 Building face database...")

    if not os.path.isdir(dataset_path):
        raise SystemExit(f"Dataset folder not found: {dataset_path} (expected {dataset_path}/<student name>/<photos>)")

    index = faiss.IndexFlatIP(EMBEDDING_DIM)
    names = []
    embeddings = []
    student_photos = {}
    all_students = []

    for student in sorted(os.listdir(dataset_path)):
        folder = os.path.join(dataset_path, student)

        if not os.path.isdir(folder):
            continue

        all_students.append(student)

        images = sorted(f for f in os.listdir(folder) if f.lower().endswith(IMAGE_EXTS))
        if images:
            student_photos[student] = os.path.join(folder, images[0])

        for img_name in images:
            img = cv2.imread(os.path.join(folder, img_name))

            if img is None:
                continue

            faces = face_app.get(img)

            if not faces:
                print(f"   ⚠️  no face found in {student}/{img_name}")
                continue

            # Use the largest face in the enrolment photo
            face = max(faces, key=lambda f: (f.bbox[2] - f.bbox[0]) * (f.bbox[3] - f.bbox[1]))
            emb = face.embedding
            embeddings.append(emb / np.linalg.norm(emb))
            names.append(student)

    if not embeddings:
        raise SystemExit(f"No faces could be enrolled from {dataset_path}")

    index.add(np.array(embeddings).astype("float32"))

    print(f"✅ Loaded {len(names)} face embeddings for {len(all_students)} students")

    return index, names, all_students, student_photos


# =========================
# FACE RECOGNITION
# =========================

def recognize_face(index, names, embedding, threshold):
    embedding = embedding / np.linalg.norm(embedding)
    embedding = embedding.astype("float32").reshape(1, -1)

    scores, ids = index.search(embedding, 1)

    if scores[0][0] > threshold:
        return names[ids[0][0]], float(scores[0][0])

    return "Unknown", float(scores[0][0])


# =========================
# PROCESS FRAME
# =========================

def process_frame(frame, face_app, index, names, threshold):
    """Return [(name, score, (x1, y1, x2, y2)), ...] for every face in the frame."""
    h, w = frame.shape[:2]
    scale = min(1.0, FRAME_WIDTH / max(h, w))

    small = cv2.resize(frame, (int(w * scale), int(h * scale))) if scale < 1 else frame
    results = []

    for face in face_app.get(small):
        name, score = recognize_face(index, names, face.embedding, threshold)
        box = tuple(int(v / scale) for v in face.bbox)
        results.append((name, score, box))

    return results


def annotate(frame, results):
    frame = frame.copy()
    for name, score, (x1, y1, x2, y2) in results:
        color = (0, 200, 0) if name != "Unknown" else (0, 0, 255)
        cv2.rectangle(frame, (x1, y1), (x2, y2), color, 2)
        label = f"{name.replace('_', ' ')} ({score:.2f})"
        (tw, th), _ = cv2.getTextSize(label, cv2.FONT_HERSHEY_SIMPLEX, 0.6, 2)
        ty = max(th + 8, y1)
        cv2.rectangle(frame, (x1, ty - th - 8), (x1 + tw + 6, ty), color, -1)
        cv2.putText(frame, label, (x1 + 3, ty - 5), cv2.FONT_HERSHEY_SIMPLEX, 0.6, (255, 255, 255), 2)
    return frame


# =========================
# VIDEO SOURCES
# =========================

class Source:
    """A webcam / stream (live), a video file, a single image or a folder of images."""

    def __init__(self, spec):
        self.spec = spec
        self.images = None
        self.cap = None

        if spec.isdigit():
            self.live = True
            self.cap = cv2.VideoCapture(int(spec))
        elif os.path.isdir(spec):
            self.live = False
            self.images = sorted(os.path.join(spec, f) for f in os.listdir(spec) if f.lower().endswith(IMAGE_EXTS))
        elif spec.lower().endswith(IMAGE_EXTS):
            self.live = False
            self.images = [spec]
        else:
            self.cap = cv2.VideoCapture(spec)
            self.live = "://" in spec  # RTSP/HTTP streams are live, files are not

        if self.cap is not None and not self.cap.isOpened():
            raise SystemExit(f"Could not open video source: {spec}")
        if self.images is not None and not self.images:
            raise SystemExit(f"No images found in: {spec}")

    def read(self, position):
        """Read a frame; `position` in [0, 1) picks the point in a recorded video / image folder."""
        if self.images is not None:
            return cv2.imread(self.images[min(int(position * len(self.images)), len(self.images) - 1)])

        if self.live:
            for _ in range(5):  # drop frames buffered while we were waiting
                self.cap.grab()
        else:
            total = int(self.cap.get(cv2.CAP_PROP_FRAME_COUNT))
            self.cap.set(cv2.CAP_PROP_POS_FRAMES, int(position * max(total - 1, 0)))

        ok, frame = self.cap.read()
        return frame if ok else None

    def release(self):
        if self.cap is not None:
            self.cap.release()


# =========================
# RANDOM SAMPLING TIMES
# =========================

def generate_random_times(duration, n):
    """n random offsets (seconds) within the lecture, in order."""
    return sorted(random.uniform(0, duration) for _ in range(n))


# =========================
# MAIN ATTENDANCE LOGIC
# =========================

def run_attendance(face_app, index, names, all_students, args):

    print("🎥 Initializing cameras...")
    sources = [Source(s) for s in args.source]
    live = any(s.live for s in sources)

    attendance_counts = {student: 0 for student in all_students}
    sample_log = []

    offsets = generate_random_times(args.duration, args.samples)
    start_time = time.time()

    if live:
        print(f"🚀 Taking {args.samples} samples at random times over {args.duration / 60:.0f} min...")
    else:
        print(f"🚀 Taking {args.samples} random samples from the recording...")

    for i, offset in enumerate(offsets):

        if live:
            while time.time() < start_time + offset:
                time.sleep(0.5)

        frame_results = []

        for cam, src in enumerate(sources):
            frame = src.read(offset / args.duration if args.duration else 0)
            if frame is None:
                print(f"   ⚠️  camera {src.spec}: no frame")
                continue

            results = process_frame(frame, face_app, index, names, args.threshold)
            frame_results.extend(name for name, _, _ in results if name != "Unknown")

            if args.snapshots:
                os.makedirs(args.snapshots, exist_ok=True)
                cv2.imwrite(os.path.join(args.snapshots, f"sample{i + 1:02d}_cam{cam}.jpg"), annotate(frame, results))

        # MULTI-CAMERA FUSION: a student counts once per sample, whichever camera saw them
        unique_names = sorted(set(frame_results))

        for name in unique_names:
            attendance_counts[name] += 1

        stamp = datetime.now().strftime("%H:%M:%S") if live else f"{offset / 60:.1f} min"
        sample_log.append((i + 1, stamp, ", ".join(n.replace("_", " ") for n in unique_names) or "-"))
        print(f"📸 Sample {i + 1}/{args.samples} ({stamp}): {sample_log[-1][2]}")

    for src in sources:
        src.release()

    return attendance_counts, sample_log


# =========================
# GENERATE REPORT
# =========================

def generate_report(attendance_counts, all_students, student_photos, sample_log, args):

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    header = ["Student", "Detected", "Attendance %", "Status", "Photo"]
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = {"A": 24, "B": 12, "C": 15, "D": 12, "E": 14}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    present_fill = PatternFill("solid", fgColor="C6EFCE")
    absent_fill = PatternFill("solid", fgColor="FFC7CE")

    for row, student in enumerate(all_students, start=2):

        count = attendance_counts.get(student, 0)
        status = "Present" if count >= args.present else "Absent"

        ws.cell(row=row, column=1, value=student.replace("_", " "))
        ws.cell(row=row, column=2, value=f"{count}/{args.samples}")
        ws.cell(row=row, column=3, value=round(100 * count / args.samples))
        ws.cell(row=row, column=4, value=status).fill = present_fill if status == "Present" else absent_fill
        for col in range(1, 5):
            ws.cell(row=row, column=col).alignment = Alignment(horizontal="center", vertical="center")

        photo_path = student_photos.get(student)

        if photo_path and os.path.exists(photo_path):
            img = ExcelImage(photo_path)
            ratio = 80 / max(img.width, img.height)  # fit in 80x80 px, keep aspect ratio
            img.width, img.height = int(img.width * ratio), int(img.height * ratio)
            ws.add_image(img, f"E{row}")
            ws.row_dimensions[row].height = 64  # points; fits the 80 px photo

    log = wb.create_sheet("Samples")
    log.append(["Sample", "Time", "Recognised"])
    for cell in log[1]:
        cell.font = Font(bold=True)
    for entry in sample_log:
        log.append(list(entry))
    log.column_dimensions["B"].width = 12
    log.column_dimensions["C"].width = 60

    wb.save(args.report)
    print(f"📊 Report saved: {args.report}")


# =========================
# MAIN ENTRY POINT
# =========================

def parse_args():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dataset", default="dataset", help="folder with one sub-folder of photos per student")
    ap.add_argument("--source", action="append",
                    help="camera index, RTSP/HTTP URL, video file, image or image folder (repeat for more cameras)")
    ap.add_argument("--duration", type=float, default=60 * 60, help="lecture length in seconds (default 3600)")
    ap.add_argument("--samples", type=int, default=10, help="random checks during the lecture (default 10)")
    ap.add_argument("--present", type=int, default=7, help="checks a student must appear in to be Present (default 7)")
    ap.add_argument("--threshold", type=float, default=0.45, help="cosine similarity needed for a match (default 0.45)")
    ap.add_argument("--report", default="attendance_report.xlsx", help="Excel report path")
    ap.add_argument("--snapshots", help="save annotated frames of every sample to this folder")
    ap.add_argument("--seed", type=int, help="random seed for reproducible sampling times")
    args = ap.parse_args()
    args.source = args.source or ["0"]
    if args.present > args.samples:
        ap.error("--present cannot be larger than --samples")
    return args


def main():
    args = parse_args()
    if args.seed is not None:
        random.seed(args.seed)

    face_app = load_model()

    index, names, all_students, student_photos = build_database(face_app, args.dataset)

    attendance_counts, sample_log = run_attendance(face_app, index, names, all_students, args)

    generate_report(attendance_counts, all_students, student_photos, sample_log, args)

    present = sum(1 for s in all_students if attendance_counts[s] >= args.present)
    print(f"✅ Attendance process completed: {present}/{len(all_students)} present")


if __name__ == "__main__":
    main()
